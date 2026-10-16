"""Parse the Access→GrowHub mapping workbook and cross-check it against both schemas."""
import openpyxl, json, re, os, sys, collections, datetime, difflib
import xml.etree.ElementTree as ET
HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = sys.argv[1] if len(sys.argv) > 1 else '/Users/christopherward/Downloads/Access_Mapping_Spreadsheet.xlsx'
SEED = '/Users/christopherward/Documents/GitHub/Nursery-GrowHub-Backend/seed/access'
OUT = os.environ.get('OUT', os.path.dirname(HERE))
G = json.load(open(os.path.join(OUT, 'growhub.json')))['tables']
L = json.load(open(os.path.join(OUT, 'nmis_frontend.json')))['erd']['tables']

# Section → GrowHub table. Not stated in the workbook; inferred from section names and checked against the fields.
SECTION_TABLE = {
    ('Source Lot Data', 'Details'): 'source_lot', ('Source Lot Data', 'Storage'): 'storage', ('Source Lot Data', 'Processing'): 'processing_activity',
    ('Source Lot Data', 'Test'): 'source_lot_test', ('Source Lot Data', 'Stratification Data'): 'source_lot_test',
    ('Source Lot Data', 'Receiving'): 'receiving_activity', ('Source Lot Data', 'Ship Ticket'): 'source_shipment', ('Source Lot Data', 'Lot Slip'): 'lot_slip',
    ('Seedling Lot', 'Request'): 'grow_request', ('Seedling Lot', 'Planning'): 'order_cycle_planning', ('Seedling Lot', 'Production'): 'order_cycle_production',
    ('Seedling Lot', 'Production Location'): 'production_location', ('Seedling Lot', 'Fulfillment - Lift Request'): 'lift_request',
    ('Seedling Lot', 'Fulfillment - Lift Summary'): 'lift_summary', ('Seedling Lot', 'Fulfillment - Pack Summary'): 'pack_summary',
    ('Seedling Lot', 'Fulfillment - Storage Summary'): 'storage_summary', ('Seedling Lot', 'Fulfillment - Thaw Summary'): 'thaw_summary',
    ('Seedling Lot', 'Inventory Progress'): 'inventory_progress', ('Seedling Lot', 'Shipping'): 'seedling_ship_ticket',
    ('Client', 'Associate'): 'client', ('Client', 'Contacts'): 'client',
    ('Species Product', 'Species Product'): 'species_product', ('Species Product', 'Seed Factor'): 'seed_factor',
    ('Species Product', 'Grow Cycle'): 'grow_cycle', ('Species Product', 'Product Inventory Cycle'): 'product_inventory_cycle',
}
JSON_SUBFIELDS = {('source_lot_test', 'Stratification Data'): 'stratificationData', ('client', 'Contacts'): 'contacts'}
EMBED_ALIAS = {'containerspec': 'containerspecifications', 'barerootspec': 'barerootspecifications'}

def clean(v):
    if v is None: return ''
    if isinstance(v, (datetime.time, datetime.datetime)): v = v.isoformat()
    return re.sub(r'\s+', ' ', str(v).replace('\xa0', ' ')).strip()
def norm(s): return re.sub(r'[^a-z0-9]', '', s.lower())
def hyperlink(v):
    m = re.match(r'=HYPERLINK\("([^"]+)"\s*,\s*"([^"]*)"\)', v or '')
    return (m[1], m[2]) if m else None

def status(v):
    s = clean(v); l = s.lower()
    if not s: return {'state': 'blank', 'text': ''}
    if l in ('yes', 'y') or l.startswith('yes,') and 'maybe' not in l and 'condition' not in l: return {'state': 'yes', 'text': s}
    if 'maybe' in l or l.endswith('?') or 'condition' in l: return {'state': 'maybe', 'text': s}
    if l.startswith('no'): return {'state': 'no', 'text': s}
    return {'state': 'other', 'text': s}

def resolve_gh(table, field, section):
    """Find the GrowHub column a workbook field refers to."""
    if not field: return None
    cols = {c['name']: c for c in G.get(table, {}).get('columns', [])}
    jf = JSON_SUBFIELDS.get((table, section))
    if jf: return {'table': table, 'column': jf, 'jsonKey': field, 'how': 'key inside JSON column'}
    f = field.strip()
    cand = []
    if '.' in f:
        a, b = f.split('.', 1); na = EMBED_ALIAS.get(norm(a), norm(a)); cand = [na + norm(b)]
    else:
        n = norm(f); cand = [n, n + 'uuid', n.replace('uuid', '') if n.endswith('uuid') else None]
    def look(cols):
        idx = {norm(c): c for c in cols}
        for c in cand:
            if c and c in idx: return idx[c]
    hit = look(cols)
    if hit: return {'table': table, 'column': hit, 'how': 'exact' if hit == f else 'normalized'}
    for t, tb in G.items():  # embedded specs sometimes live on a sibling table
        hit = look({c['name']: c for c in tb['columns']})
        if hit and '.' in f: return {'table': t, 'column': hit, 'how': 'found on another table'}
    return None

def suggest_gh(table, field):
    if not field or not table: return []
    f = norm(field.split('.')[-1]).replace('uuid', '')
    out = []
    cols = [c['name'] for c in G.get(table, {}).get('columns', []) if not c['name'].startswith('metadata')]
    nc = {norm(c).replace('uuid', ''): c for c in cols}
    for m in difflib.get_close_matches(f, list(nc), n=3, cutoff=0.6): out.append({'table': table, 'column': nc[m]})
    for c in cols:
        n = norm(c).replace('uuid', '')
        if f and (f in n or n in f) and len(min(f, n, key=len)) >= 4 and not any(o['column'] == c for o in out): out.append({'table': table, 'column': c})
    for t, tb in G.items():  # same name on another table
        if t == table: continue
        for c in tb['columns']:
            if norm(c['name']).replace('uuid', '') == f: out.append({'table': t, 'column': c['name']})
    return out[:4]

# legacy columns: frontend extract + the four nurseries' XML exports
legacy_cols = {t: set(v['columns']) for t, v in L.items()}
xml_cols = collections.defaultdict(set)
for nur in ('jhsn', 'lpn', 'pn', 'tn'):
    d = os.path.join(SEED, nur)
    if not os.path.isdir(d): continue
    for fn in os.listdir(d):
        if not fn.startswith('NMIS_') or not fn.endswith('.xml'): continue
        t = fn[:-4]; n = 0
        for ev, el in ET.iterparse(os.path.join(d, fn), events=('end',)):
            if el.tag == t:
                for ch in el: xml_cols[t].add(ch.tag)
                n += 1; el.clear()
                if n > 3000: break
for t, cs in xml_cols.items(): legacy_cols.setdefault(t, set()).update(cs)
def legacy_table(name):
    n = clean(name).upper()
    for c in (n, n + 'S', n.rstrip('S')):
        if c in legacy_cols: return c
    return None
def resolve_legacy(tables, field):
    out = []
    tnames = [legacy_table(t) or clean(t) for t in re.split(r'\s*(?:\+|/|,)\s*', tables or '') if clean(t)]
    toks = re.findall(r'[A-Za-z_][A-Za-z0-9_]*', field or '')
    for tok in toks:
        tok2 = tok.split('.')[-1]
        for t in tnames:
            cs = legacy_cols.get(t, set())
            hit = next((c for c in cs if c.upper() == tok2.upper()), None)
            if hit: out.append({'table': t, 'column': hit}); break
        else:
            if '.' in tok and legacy_table(tok.split('.')[0]):
                t = legacy_table(tok.split('.')[0]); hit = next((c for c in legacy_cols[t] if c.upper() == tok2.upper()), None)
                if hit: out.append({'table': t, 'column': hit})
            elif not any(o for o in out if o.get('_tok') == tok):
                # joined-query names such as SPECIES_PRODUCT_SPEC_MIN = NMIS_SPECIES_PRODUCTS.SPEC_MIN
                for lt, cs in legacy_cols.items():
                    stem = re.sub(r'(IE)?S$', '', lt.replace('NMIS_', '')) + '_'
                    if tok2.upper().startswith(stem) and len(tok2) > len(stem):
                        rest = tok2[len(stem):]
                        hit = next((c for c in cs if c.upper() == rest.upper()), None)
                        if hit: out.append({'table': lt, 'column': hit, 'via': tok2}); break
    return tnames, out

wbf = openpyxl.load_workbook(XLSX)            # formulas (hyperlinks)
wbv = openpyxl.load_workbook(XLSX, data_only=True)  # cached values (VLOOKUPs)

# ---- enum tables (SourceENUM) ----
enums = []
ws = wbf['SourceENUM']
for c in range(1, ws.max_column + 1):
    if clean(ws.cell(1, c).value).lower() != 'access value': continue
    name_a, name_g = clean(ws.cell(2, c).value), clean(ws.cell(2, c + 1).value)
    vals = []
    for r in range(3, ws.max_row + 1):
        a, g = ws.cell(r, c).value, ws.cell(r, c + 1).value
        if a is None and g is None: continue
        vals.append({'access': clean(a), 'growhub': clean(g)})
    col = openpyxl.utils.get_column_letter(c)
    enums.append({'id': col, 'access': name_a or '(unnamed)', 'growhub': name_g, 'values': vals})
enum_by_col = {e['id']: e for e in enums}

# ---- field mapping sheets ----
domains = []
for sheet in ('Source Lot Data', 'Seedling Lot', 'Client', 'Species Product'):
    ws = wbf[sheet]
    hdr = {clean(ws.cell(1, c).value).lower(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    def col(*names):
        for n in names:
            for h, c in hdr.items():
                if h.startswith(n): return c
    C = {'section': col('section'), 'gh': col('growhub field'), 'at': col('access table'), 'af': col('access field'),
         'approved': col('approved'), 'validated': col('validated'), 'enum': col('enum'), 'enummap': col('enum mapping'),
         'notes': col('notes'), 'final': col('final approval')}
    rows = []; last_sec = ''
    for r in range(2, ws.max_row + 1):
        get = lambda k: ws.cell(r, C[k]).value if C[k] else None
        sec = clean(get('section')); gh = clean(get('gh')); at = clean(get('at')); af = clean(get('af'))
        if not (gh or at or af): continue
        if sec: last_sec = sec
        else: sec = last_sec
        if sheet == 'Species Product' and r > 51: continue  # worked example block below the table
        table = SECTION_TABLE.get((sheet, sec))
        em_raw = get('enummap'); em = clean(em_raw); link = hyperlink(em_raw if isinstance(em_raw, str) else '')
        enum_ref = None; ext = None; em_text = em
        if link:
            em_text = ''
            if link[0].startswith('#'):
                m = re.search(r"!\$?([A-Z]+)\$?\d", link[0]); enum_ref = m[1] if m and m[1] in enum_by_col else None
            else: ext = link[0]
        notes_raw = get('notes'); nl = hyperlink(notes_raw if isinstance(notes_raw, str) else '')
        notes = clean(notes_raw)
        if nl: notes = ''; ext = ext or nl[0]
        elif notes.startswith('http'): ext = ext or notes; notes = ''
        fill = ws.cell(r, 1).fill.fgColor.rgb if ws.cell(r, 1).fill else None
        tnames, lres = resolve_legacy(at, af)
        ghres = resolve_gh(table, gh, sec) if table else None
        rows.append({'row': r, 'section': sec, 'table': table, 'field': gh, 'accessTables': at, 'accessField': af,
                     'approved': status(get('approved')), 'validated': status(get('validated')), 'final': status(get('final')),
                     'enum': clean(get('enum')).lower() == 'yes', 'enumMapping': em_text, 'enumRef': enum_ref, 'link': ext,
                     'notes': notes, 'flagged': bool(fill and fill not in ('00000000',)),
                     'gh': ghres, 'legacy': lres, 'legacyTables': tnames,
                     'suggest': suggest_gh(table, gh) if table and gh and not ghres else []})
    secs = collections.OrderedDict()
    for x in rows: secs.setdefault(x['section'], []).append(x)
    domains.append({'sheet': sheet, 'sections': [{'name': k, 'table': SECTION_TABLE.get((sheet, k)), 'rows': v} for k, v in secs.items()]})

# ---- coverage: GrowHub columns of mapped tables not in the workbook ----
AUTO = re.compile(r'^(uuid|metadata.*|legacyData|legacyId|accessLegacyId)$')
mapped = collections.defaultdict(set)
for d in domains:
    for s in d['sections']:
        for x in s['rows']:
            if x['gh']: mapped[x['gh']['table']].add(x['gh']['column'])
coverage = {}
for t in sorted({s['table'] for d in domains for s in d['sections'] if s['table']}):
    cols = [c['name'] for c in G.get(t, {}).get('columns', [])]
    fks = {f['column'] for f in G.get(t, {}).get('fks', [])}
    coverage[t] = {'mapped': sorted(mapped[t]), 'unmapped': [c for c in cols if c not in mapped[t] and not AUTO.match(c)],
                   'unmappedFks': [c for c in cols if c in fks and c not in mapped[t] and not AUTO.match(c)]}

# ---- activity codes ----
ws = wbv['Activity Codes Worksheet']
oracle = []; access = []
for r in range(4, ws.max_row + 1):
    v = lambda c: clean(ws.cell(r, c).value)
    if v(2): oracle.append({'code': v(2), 'name': v(3), 'ghParent': v(4), 'ghCode': v(5), 'kind': v(6), 'balance': v(7), 'balanceCodes': v(8)})
    if v(10) and v(11): access.append({'nursery': v(10), 'code': v(11), 'name': v(12), 'parentCode': v(13), 'contractItem': v(14), 'billOrder': v(15),
                                       'balance': v(16), 'codeNo': v(17), 'note': v(18), 'module': v(19), 'kind': v(20), 'ghParent': v(21), 'ghCode': v(22)})
# ---- grow units / box types ----
ws = wbv['Grow Units']
units = []; desired = []; mismatched = []
for r in range(3, ws.max_row + 1):
    v = lambda c: ws.cell(r, c).value
    if v(1) and v(2): units.append({'nursery': clean(v(1)), 'code': clean(v(2)), 'type': clean(v(3)), 'multiplier': v(4), 'description': clean(v(5)),
                                    'cellDims': clean(v(6)), 'cellsPerContainer': v(7), 'areaPerContainer': v(15), 'ghCode': clean(v(17)), 'notes': clean(v(18)), 'approved': status(v(19))})
    if v(24): desired.append({'current': clean(v(24)), 'desired': clean(v(25)), 'description': clean(v(26)), 'stockType': clean(v(27)), 'cellDiameter': v(28),
                              'cellDepth': v(29), 'cellDims': clean(v(30)), 'cellsPerContainer': v(31), 'areaPerContainer': v(32), 'width': v(33), 'length': v(34),
                              'height': v(35), 'drillRows': v(36), 'historical': clean(v(37))})
    if v(41) is not None and v(42): mismatched.append({'value': clean(v(41)), 'nursery': clean(v(42)), 'firstProduct': clean(v(43))})
ws = wbv['Ship|Storage Box Type']
boxes = [{'nursery': clean(ws.cell(r, 1).value), 'code': clean(ws.cell(r, 2).value), 'type': clean(ws.cell(r, 3).value), 'multiplier': ws.cell(r, 4).value,
          'description': clean(ws.cell(r, 5).value), 'notes': clean(ws.cell(r, 6).value), 'cellsPerContainer': ws.cell(r, 7).value}
         for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value]

out = {'source': os.path.basename(XLSX), 'modified': datetime.datetime.fromtimestamp(os.path.getmtime(XLSX)).strftime('%Y-%m-%d %H:%M'),
       'domains': domains, 'enums': enums, 'coverage': coverage, 'sectionTables': {'%s|%s' % k: v for k, v in SECTION_TABLE.items()},
       'activityCodes': {'oracle': oracle, 'access': access}, 'growUnits': {'units': units, 'desired': desired, 'mismatched': mismatched},
       'boxTypes': boxes, 'legacyXmlColumns': {t: sorted(c) for t, c in xml_cols.items()}}
json.dump(out, open(os.path.join(OUT, 'mapping.json'), 'w'), indent=1, default=str)

rows = [x for d in domains for s in d['sections'] for x in s['rows']]
print(len(rows), 'mapping rows;', sum(1 for x in rows if x['gh']), 'resolved to a GrowHub column;', sum(1 for x in rows if x['legacy']), 'resolved to a legacy column')
for x in rows:
    if x['field'] and x['table'] and not x['gh']: print('  not found', x['table'], x['field'], '->', [o['table'] + '.' + o['column'] for o in x['suggest']])
print('Access field given but not found:', [(x['section'], x['accessTables'], x['accessField']) for x in rows if x['accessField'] and not x['legacy']])
print('status:', collections.Counter(x['approved']['state'] for x in rows))
print('activity codes', len(oracle), len(access), 'access mapped', sum(1 for a in access if a['ghCode'] or a['ghParent']))
print('units', len(units), 'desired', len(desired), 'mismatched', len(mismatched), 'boxes', len(boxes), 'enums', [(e['access'], len(e['values'])) for e in enums])
print('xml tables', {t: len(c) for t, c in xml_cols.items()})
